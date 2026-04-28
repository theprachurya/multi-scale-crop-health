import openpyxl
wb = openpyxl.load_workbook(r'Supporting Documents\Sprint Retrospective (1).xlsx')
ws = wb.active

for r in range(1, ws.max_row+1):
    for c in range(1, ws.max_column+1):
        val = ws.cell(r, c).value
        if val:
            print(f"=== R{r}C{c} ===")
            print(val)
            print()
